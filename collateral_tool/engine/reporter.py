"""
Generation du rapport Excel multi-onglets avec formatage professionnel.
"""

from pathlib import Path
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ── Palette de couleurs ───────────────────────────────────────────────────────

POLICE      = "Arial"
TAILLE      = 10
HEADER_BG   = "1F3864"
GRIS_LIGNE  = "F2F2F2"
BLANC_LIGNE = "FFFFFF"
TOTAL_BG    = "D6E4F0"
KPI_BG      = "EBF3FB"

STATUT_STYLE: dict[str, tuple[str, str]] = {
    "SUBSTITUT TROUVE":     ("C6EFCE", "276221"),
    "COUVERTURE PARTIELLE": ("FFEB9C", "9C5700"),
    "AUCUN SUBSTITUT":      ("FFC7CE", "9C0006"),
}

ALERTE_STYLE: dict[str, tuple[str, str]] = {
    "URGENT":    ("FFC7CE", "9C0006"),
    "ATTENTION": ("FFEB9C", "9C5700"),
    "OK":        ("C6EFCE", "276221"),
}

FMT_MONETAIRE = "#,##0"
FMT_PCT       = "0.0%"
FMT_DATE      = "DD/MM/YYYY"


# ── Helpers ───────────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _font(bold: bool = False, color: str = "000000") -> Font:
    return Font(name=POLICE, size=TAILLE, bold=bold, color=color)


def _categorie_alerte(jours: int) -> str:
    if jours < 7:
        return "URGENT"
    if jours < 30:
        return "ATTENTION"
    return "OK"


# ── Builder de feuille ────────────────────────────────────────────────────────

class _FeuilleBuilder:
    def __init__(self, ws, colonnes: list[str], formats_cols: dict[int, str]):
        self.ws           = ws
        self.colonnes     = colonnes
        self.n_cols       = len(colonnes)
        self.formats_cols = formats_cols
        self.row          = 1

        ws.sheet_view.showGridLines = False
        self._ecrire_header()

    def _ecrire_header(self) -> None:
        self.ws.row_dimensions[1].height = 28
        for ci, nom in enumerate(self.colonnes, 1):
            c = self.ws.cell(row=1, column=ci, value=nom)
            c.fill      = _fill(HEADER_BG)
            c.font      = _font(bold=True, color="FFFFFF")
            c.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        self.ws.freeze_panes = self.ws.cell(row=2, column=1)
        self.row = 2

    def ajouter_ligne(
        self,
        valeurs: list,
        surcharges: dict[int, tuple[str, str]] | None = None,
    ) -> None:
        bg = GRIS_LIGNE if (self.row % 2 == 0) else BLANC_LIGNE
        for ci, val in enumerate(valeurs, 1):
            c = self.ws.cell(row=self.row, column=ci, value=val)
            c.fill      = _fill(bg)
            c.font      = _font()
            c.alignment = Alignment(
                horizontal="right" if ci in self.formats_cols else "left",
                vertical="center",
            )
            if ci in self.formats_cols:
                c.number_format = self.formats_cols[ci]

        if surcharges:
            for ci, (bg_h, fg_h) in surcharges.items():
                c = self.ws.cell(row=self.row, column=ci)
                c.fill      = _fill(bg_h)
                c.font      = _font(bold=True, color=fg_h)
                c.alignment = Alignment(horizontal="center", vertical="center")

        self.row += 1

    def ajouter_total(self, totaux: dict[int, float]) -> None:
        for ci in range(1, self.n_cols + 1):
            c = self.ws.cell(row=self.row, column=ci)
            c.fill      = _fill(TOTAL_BG)
            c.font      = _font(bold=True)
            c.alignment = Alignment(vertical="center")

        self.ws.cell(row=self.row, column=1).value = "TOTAL"

        for ci, val in totaux.items():
            c = self.ws.cell(row=self.row, column=ci)
            c.value         = val
            c.fill          = _fill(TOTAL_BG)
            c.font          = _font(bold=True)
            c.alignment     = Alignment(horizontal="right", vertical="center")
            if ci in self.formats_cols:
                c.number_format = self.formats_cols[ci]

        self.row += 1

    def ajuster_largeurs(self) -> None:
        for ci, nom in enumerate(self.colonnes, 1):
            col_letter = get_column_letter(ci)
            max_len    = len(nom)
            for rows in self.ws.iter_rows(
                min_row=2, max_row=self.row - 1, min_col=ci, max_col=ci
            ):
                for cell in rows:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
            self.ws.column_dimensions[col_letter].width = min(max_len + 4, 50)


# ── Onglet 1 : Tableau de Bord ────────────────────────────────────────────────

def _ecrire_tableau_de_bord(
    ws,
    df_poste: pd.DataFrame,
    df_dispo: pd.DataFrame,
    df_subs: pd.DataFrame,
) -> None:
    ws.sheet_view.showGridLines = False

    def _kpi_bloc(row: int, col: int, label: str, valeur, fmt: str = "") -> None:
        c_label = ws.cell(row=row, column=col, value=label)
        c_label.fill      = _fill(HEADER_BG)
        c_label.font      = _font(bold=True, color="FFFFFF")
        c_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        c_val = ws.cell(row=row + 1, column=col, value=valeur)
        c_val.fill          = _fill(KPI_BG)
        c_val.font          = Font(name=POLICE, size=14, bold=True, color="1F3864")
        c_val.alignment     = Alignment(horizontal="center", vertical="center")
        if fmt:
            c_val.number_format = fmt

        ws.row_dimensions[row].height     = 18
        ws.row_dimensions[row + 1].height = 32
        ws.column_dimensions[get_column_letter(col)].width = 26

    aujourd_hui = date.today()

    # KPIs colonne 1 : exposition totale
    ve_totale   = df_poste["Valeur_Eligible"].sum()
    nb_titres   = len(df_poste)
    _kpi_bloc(2, 2, "Exposition totale (VE)", round(ve_totale), FMT_MONETAIRE)
    _kpi_bloc(6, 2, "Nombre de titres postes", nb_titres)

    # KPIs colonne 2 : alertes
    nb_alertes = len(df_poste[df_poste["Jours_Restants"] < 30])
    ve_alertes = df_poste[df_poste["Jours_Restants"] < 30]["Valeur_Eligible"].sum()
    _kpi_bloc(2, 4, "Titres en alerte (< 30j)", nb_alertes)
    _kpi_bloc(6, 4, "Exposition en alerte (VE)", round(ve_alertes), FMT_MONETAIRE)

    # KPIs colonne 3 : substitutions
    if not df_subs.empty:
        nb_trouves  = (df_subs["Statut"] == "SUBSTITUT TROUVE").sum()
        nb_partiels = (df_subs["Statut"] == "COUVERTURE PARTIELLE").sum()
        nb_aucun    = (df_subs["Statut"] == "AUCUN SUBSTITUT").sum()
        taux = nb_trouves / len(df_subs) if len(df_subs) > 0 else 0
    else:
        nb_trouves = nb_partiels = nb_aucun = 0
        taux = 0.0

    _kpi_bloc(2, 6, "Substituts trouves", nb_trouves)
    _kpi_bloc(6, 6, "Taux de couverture", taux, "0%")

    # KPIs colonne 4 : statuts detailles
    _kpi_bloc(2, 8, "Couverture partielle", nb_partiels)
    _kpi_bloc(6, 8, "Aucun substitut", nb_aucun)

    # Colorer les cellules critiques
    if nb_aucun > 0:
        c = ws.cell(row=7, column=8)
        c.fill = _fill("FFC7CE")
        c.font = Font(name=POLICE, size=14, bold=True, color="9C0006")
    if nb_partiels > 0:
        c = ws.cell(row=3, column=8)
        c.fill = _fill("FFEB9C")
        c.font = Font(name=POLICE, size=14, bold=True, color="9C5700")
    if nb_trouves > 0:
        c = ws.cell(row=3, column=6)
        c.fill = _fill("C6EFCE")
        c.font = Font(name=POLICE, size=14, bold=True, color="276221")

    # Repartition par type (tableau bas)
    ws.row_dimensions[10].height = 6
    row_titre = 11
    c = ws.cell(row=row_titre, column=2, value="Repartition du portefeuille poste par type d'actif")
    c.fill = _fill(HEADER_BG)
    c.font = _font(bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(f"B{row_titre}:I{row_titre}")

    row_data = row_titre + 1
    repartition = df_poste.groupby("Type")["Valeur_Eligible"].sum().sort_values(ascending=False)
    alternance  = [BLANC_LIGNE, GRIS_LIGNE]
    for i, (type_actif, ve) in enumerate(repartition.items()):
        bg   = alternance[i % 2]
        pct  = ve / ve_totale if ve_totale > 0 else 0
        for ci in [2, 3, 4, 5]:
            ws.cell(row=row_data, column=ci).fill = _fill(bg)
        c = ws.cell(row=row_data, column=2, value=type_actif)
        c.fill = _fill(bg)
        c.font = _font()
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c = ws.cell(row=row_data, column=4, value=round(ve))
        c.fill          = _fill(bg)
        c.font          = _font()
        c.number_format = FMT_MONETAIRE
        c.alignment     = Alignment(horizontal="right", vertical="center")
        c = ws.cell(row=row_data, column=5, value=pct)
        c.fill          = _fill(bg)
        c.font          = _font()
        c.number_format = "0.0%"
        c.alignment     = Alignment(horizontal="right", vertical="center")
        row_data += 1

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["C"].width = 4


# ── Onglet 2 : Alertes & Substitutions ───────────────────────────────────────

def _ecrire_alertes_substitutions(ws, df: pd.DataFrame) -> None:
    colonnes = [
        "Contrepartie", "ISIN Sortant", "Titre Sortant", "Maturite",
        "VE Sortant", "ISIN Substitut", "Titre Substitut",
        "VE Substitut", "Score CTD", "Statut",
    ]
    formats = {4: FMT_DATE, 5: FMT_MONETAIRE, 8: FMT_MONETAIRE}
    builder = _FeuilleBuilder(ws, colonnes, formats)

    for _, row in df.iterrows():
        statut = row["Statut"]
        ve_sub = row["Valeur_Eligible_Substitut"]
        score  = row["Score_CTD"]

        valeurs = [
            row["Contrepartie"],
            row["ISIN_Sortant"],
            row["Titre_Sortant"],
            row["Maturite"].date() if hasattr(row["Maturite"], "date") else row["Maturite"],
            row["Valeur_Eligible_Sortant"],
            row["ISIN_Substitut"] or None,
            row["Titre_Substitut"] or None,
            int(ve_sub) if ve_sub else None,
            int(score) if pd.notna(score) else None,
            statut,
        ]

        surcharges = {}
        if statut in STATUT_STYLE:
            surcharges[10] = STATUT_STYLE[statut]

        builder.ajouter_ligne(valeurs, surcharges)

    totaux = {
        5: int(df["Valeur_Eligible_Sortant"].sum()),
        8: int(df["Valeur_Eligible_Substitut"].sum()),
    }
    builder.ajouter_total(totaux)
    builder.ajuster_largeurs()


# ── Onglet 3 : Portefeuille Poste ────────────────────────────────────────────

def _ecrire_portefeuille_poste(ws, df: pd.DataFrame) -> None:
    colonnes = [
        "ISIN", "Nom", "Type", "Contrepartie", "Valeur Nominale",
        "Prix (%)", "Haircut (%)", "Maturite", "Rating",
        "Valeur Marche", "Valeur Eligible", "Jours Restants", "Alerte",
    ]
    formats = {
        5:  FMT_MONETAIRE,
        6:  FMT_PCT,
        7:  FMT_PCT,
        8:  FMT_DATE,
        10: FMT_MONETAIRE,
        11: FMT_MONETAIRE,
    }
    builder = _FeuilleBuilder(ws, colonnes, formats)

    for _, row in df.iterrows():
        jours   = int(row["Jours_Restants"])
        cat     = _categorie_alerte(jours)
        valeurs = [
            row["ISIN"],
            row["Nom"],
            row["Type"],
            row["Contrepartie"],
            int(row["Valeur_Nominale"]),
            row["Prix_Pct"] / 100,
            row["Haircut_Pct"] / 100,
            row["Maturite"].date() if hasattr(row["Maturite"], "date") else row["Maturite"],
            row["Rating"],
            round(row["Valeur_Marche"]),
            round(row["Valeur_Eligible"]),
            jours,
            cat,
        ]
        builder.ajouter_ligne(valeurs, surcharges={13: ALERTE_STYLE[cat]})

    totaux = {
        5:  int(df["Valeur_Nominale"].sum()),
        10: round(df["Valeur_Marche"].sum()),
        11: round(df["Valeur_Eligible"].sum()),
    }
    builder.ajouter_total(totaux)
    builder.ajuster_largeurs()


# ── Onglet 4 : Disponible ─────────────────────────────────────────────────────

def _ecrire_disponible(ws, df: pd.DataFrame) -> None:
    df_tri = df.sort_values("Score_CTD").reset_index(drop=True)

    colonnes = [
        "ISIN", "Nom", "Type", "Valeur Nominale", "Prix (%)",
        "Haircut (%)", "Maturite", "Rating", "Eligibilite",
        "Valeur Marche", "Valeur Eligible", "Score CTD",
    ]
    formats = {
        4:  FMT_MONETAIRE,
        5:  FMT_PCT,
        6:  FMT_PCT,
        7:  FMT_DATE,
        10: FMT_MONETAIRE,
        11: FMT_MONETAIRE,
    }
    builder = _FeuilleBuilder(ws, colonnes, formats)

    for _, row in df_tri.iterrows():
        maturite = row["Maturite"]
        valeurs  = [
            row["ISIN"],
            row["Nom"],
            row["Type"],
            int(row["Valeur_Nominale"]),
            row["Prix_Pct"] / 100,
            row["Haircut_Pct"] / 100,
            maturite.date() if hasattr(maturite, "date") else maturite,
            row["Rating"],
            row["Eligibilite"],
            round(row["Valeur_Marche"]),
            round(row["Valeur_Eligible"]),
            int(row["Score_CTD"]),
        ]
        builder.ajouter_ligne(valeurs)

    totaux = {
        4:  int(df_tri["Valeur_Nominale"].sum()),
        10: round(df_tri["Valeur_Marche"].sum()),
        11: round(df_tri["Valeur_Eligible"].sum()),
    }
    builder.ajouter_total(totaux)
    builder.ajuster_largeurs()


# ── Point d'entree public ─────────────────────────────────────────────────────

def generer_rapport(
    df_poste: pd.DataFrame,
    df_dispo: pd.DataFrame,
    df_substitutions: pd.DataFrame,
    chemin_sortie: Path,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Tableau de Bord")
    _ecrire_tableau_de_bord(ws1, df_poste, df_dispo, df_substitutions)

    ws2 = wb.create_sheet("Alertes & Substitutions")
    if not df_substitutions.empty:
        _ecrire_alertes_substitutions(ws2, df_substitutions)

    ws3 = wb.create_sheet("Portefeuille Poste")
    _ecrire_portefeuille_poste(ws3, df_poste)

    ws4 = wb.create_sheet("Disponible")
    _ecrire_disponible(ws4, df_dispo)

    wb.save(chemin_sortie)
